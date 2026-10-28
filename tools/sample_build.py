# -*- coding: utf-8 -*-
"""格式冒烟测试：按《试卷识别规则.md》生成一份最小试卷。
用法：dist/试卷识别.app/Contents/MacOS/ExamOCR --runpython tools/sample_build.py

只为验证版式硬指标：A4、四边 2.2cm、装订线 0、中文宋体 / 英文 Times New Roman、
红色答案、紫色（#800080）插图占位（1 行 1 列表格 + trHeight exact）。不追求内容质量。
"""
from docx import Document
from docx.shared import Pt, Cm, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.table import WD_TABLE_ALIGNMENT
from docx.enum.table import WD_ALIGN_VERTICAL
from docx.oxml.ns import qn
from docx.oxml import OxmlElement

BLACK = RGBColor(0x00, 0x00, 0x00)
RED = RGBColor(0xFF, 0x00, 0x00)
PURPLE = RGBColor(0x80, 0x00, 0x80)


def style_run(run, size=12, color=BLACK, bold=False):
    run.font.size = Pt(size)
    run.font.bold = bold
    run.font.color.rgb = color
    run.font.name = "Times New Roman"
    run._element.rPr.rFonts.set(qn("w:eastAsia"), "宋体")


def para(doc, text, size=12, color=BLACK, bold=False,
         align=WD_ALIGN_PARAGRAPH.LEFT, indent=True, space_after=0):
    p = doc.add_paragraph()
    p.alignment = align
    p.paragraph_format.space_before = Pt(0)
    p.paragraph_format.space_after = Pt(space_after)
    p.paragraph_format.line_spacing = 1.15
    if indent:
        p.paragraph_format.first_line_indent = Pt(24)
    style_run(p.add_run(text), size=size, color=color, bold=bold)
    return p


def figure_placeholder(doc, note="此处原为插图：一张折线图"):
    """规则里的硬要求：1 行 1 列表格模拟插图框，列宽 14cm、行高 7.5cm、紫字居中。"""
    table = doc.add_table(rows=1, cols=1)
    table.alignment = WD_TABLE_ALIGNMENT.CENTER
    table.autofit = False
    cell = table.cell(0, 0)
    cell.width = Cm(14)

    tr = table.rows[0]._tr
    trPr = tr.get_or_add_trPr()
    h = OxmlElement("w:trHeight")
    h.set(qn("w:val"), str(int(7.5 * 567)))   # 1cm ≈ 567 twips
    h.set(qn("w:hRule"), "exact")
    trPr.append(h)

    cell.vertical_alignment = WD_ALIGN_VERTICAL.CENTER
    p = cell.paragraphs[0]
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    style_run(p.add_run(note), size=10.5, color=PURPLE)


doc = Document()

sec = doc.sections[0]
sec.page_width = Cm(21.0)
sec.page_height = Cm(29.7)
sec.top_margin = Cm(2.2)
sec.bottom_margin = Cm(2.2)
sec.left_margin = Cm(2.2)
sec.right_margin = Cm(2.2)
sec.gutter = Cm(0)

# 默认样式也钉成宋体，避免 AI 只改 run 导致空段落回落到 Calibri
normal = doc.styles["Normal"]
normal.font.name = "Times New Roman"
normal.font.size = Pt(12)
normal.element.rPr.rFonts.set(qn("w:eastAsia"), "宋体")

para(doc, "2026 届高三语文第 3 套", size=16, bold=True,
     align=WD_ALIGN_PARAGRAPH.CENTER, indent=False)
para(doc, "本试卷共 8 页，满分 150 分，考试用时 150 分钟。", size=10.5,
     align=WD_ALIGN_PARAGRAPH.CENTER, indent=False, space_after=12)

para(doc, "一、现代文阅读（35 分）", size=14, bold=True, indent=False, space_after=6)
para(doc, "（一）现代文阅读Ⅰ（本题共 5 小题，19 分）")
para(doc, "阅读下面的文字，完成 1～5 题。")
para(doc, "材料一：数字技术正在重塑乡村的日常秩序。__1__ 处的空缺，恰恰说明这一过程远未完成。")
para(doc, "1．下列对材料相关内容的理解和分析，不正确的一项是（　　）")
for opt in ["A．乡村数字化始于基础设施的铺开", "B．技术下沉改变了信息获取方式",
            "C．传统手艺因此彻底消失", "D．代际差异仍然明显"]:
    para(doc, opt, indent=False)
para(doc, "【答案】C", color=RED, indent=False)
para(doc, "【解析】原文只说“部分手艺面临失传”，并未说“彻底消失”，属过度推断。",
     color=RED, indent=False, space_after=12)

figure_placeholder(doc)

doc.save("试卷冒烟测试.docx")
print("saved 试卷冒烟测试.docx")
